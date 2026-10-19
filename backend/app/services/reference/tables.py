"""Read a CSV or Excel sheet into dicts with our own column names.

Official sheets use many headings for the same thing ("CGHS Code", "Code",
"Procedure Code"...). ALIASES maps each of our columns to the headings we
accept. The header row does not have to be the first row: the first row with
at least two recognised headings is used, so title rows above the table are fine.
"""

import csv
import gzip
from collections.abc import Iterator
from pathlib import Path

from app.services.reference.text import header_key


class TableError(Exception):
    pass


def _rows_from_file(path: Path, sheet: str | None) -> Iterator[list[object]]:
    suffix = "".join(path.suffixes[-2:]).lower() if path.name.endswith(".gz") else path.suffix.lower()
    if suffix in {".csv", ".csv.gz"}:
        opener = gzip.open if suffix.endswith(".gz") else open
        with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:  # type: ignore[operator]
            yield from csv.reader(handle)
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = book[sheet] if sheet else book.worksheets[0]
        except KeyError:
            raise TableError(f"Sheet '{sheet}' not found. Sheets: {', '.join(book.sheetnames)}")
        for row in ws.iter_rows(values_only=True):
            yield list(row)
        book.close()
    else:
        raise TableError(f"Unsupported file type '{path.suffix}'. Use .csv or .xlsx")


def read_table(path: Path, aliases: dict[str, list[str]], required: list[str], sheet: str | None = None,
               ) -> list[dict[str, object]]:
    lookup = {header_key(a): col for col, names in aliases.items() for a in [col, *names]}
    rows = _rows_from_file(path, sheet)
    mapping: dict[int, str] | None = None
    for raw in rows:
        found = {i: lookup[header_key(v)] for i, v in enumerate(raw) if header_key(v) in lookup}
        if len(set(found.values())) >= 2:
            mapping = found
            break
    if mapping is None:
        raise TableError("Could not find a header row. Expected columns like: " + ", ".join(aliases))
    missing = [c for c in required if c not in mapping.values()]
    if missing:
        raise TableError("Missing required column(s): " + ", ".join(missing))

    out: list[dict[str, object]] = []
    for raw in rows:
        record = {col: (raw[i] if i < len(raw) else None) for i, col in mapping.items()}
        if all(v is None or str(v).strip() == "" for v in record.values()):
            continue
        out.append(record)
    return out
